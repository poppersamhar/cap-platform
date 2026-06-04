"""文档解析 — 支持 PDF / Word / PPT / TXT / Markdown"""

import io
import logging
from typing import BinaryIO

logger = logging.getLogger("cap.knowledge")


def parse_document(file_bytes: bytes, filename: str) -> str:
    """根据文件扩展名选择解析器，返回纯文本"""
    name = filename.lower()

    if name.endswith(".pdf"):
        return _parse_pdf(file_bytes)
    elif name.endswith(".docx"):
        return _parse_docx(file_bytes)
    elif name.endswith(".pptx"):
        return _parse_pptx(file_bytes)
    elif name.endswith(".txt") or name.endswith(".md"):
        return file_bytes.decode("utf-8", errors="ignore")
    else:
        # fallback：尝试按文本解码
        return file_bytes.decode("utf-8", errors="ignore")


def _parse_pdf(file_bytes: bytes) -> str:
    """PDF 解析"""
    try:
        import pdfplumber
    except ImportError:
        logger.warning("pdfplumber not installed, skipping PDF text extraction")
        return ""

    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def _parse_docx(file_bytes: bytes) -> str:
    """Word 解析"""
    try:
        from docx import Document
    except ImportError:
        logger.warning("python-docx not installed, skipping DOCX text extraction")
        return ""

    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def _parse_pptx(file_bytes: bytes) -> str:
    """PPT 解析"""
    try:
        from pptx import Presentation
    except ImportError:
        logger.warning("python-pptx not installed, skipping PPTX text extraction")
        return ""

    prs = Presentation(io.BytesIO(file_bytes))
    text_parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text_parts.append(shape.text.strip())
    return "\n".join(text_parts)


def chunk_text(text: str, chunk_size: int = 300, overlap: int = 50) -> list[str]:
    """将长文本切分为语义块

    策略：
    1. 先按空行切分为段落
    2. 段落内如果超过 chunk_size，再按句子切分
    3. 相邻块之间有 overlap 重叠，保证语义连贯
    """
    # 清理文本
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # 按空行拆分为段落
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    if not paragraphs:
        return []

    chunks = []
    current_chunk = ""

    for para in paragraphs:
        # 单个段落超过 chunk_size，需要句子级切分
        if len(para) > chunk_size:
            sentences = _split_sentences(para)
            for sent in sentences:
                if len(current_chunk) + len(sent) + 1 <= chunk_size:
                    current_chunk += (" " if current_chunk else "") + sent
                else:
                    if current_chunk:
                        chunks.append(current_chunk)
                    current_chunk = sent
        else:
            if len(current_chunk) + len(para) + 2 <= chunk_size:
                current_chunk += ("\n\n" if current_chunk else "") + para
            else:
                if current_chunk:
                    chunks.append(current_chunk)
                current_chunk = para

    if current_chunk:
        chunks.append(current_chunk)

    # 应用 overlap
    if overlap > 0 and len(chunks) > 1:
        overlapped = []
        for i, chunk in enumerate(chunks):
            if i == 0:
                overlapped.append(chunk)
            else:
                prev_tail = chunks[i - 1][-overlap:]
                overlapped.append(prev_tail + "\n" + chunk)
        return overlapped

    return chunks


def _split_sentences(text: str) -> list[str]:
    """按中文/英文句号、问号、感叹号切分句子"""
    import re
    # 匹配中文和英文句子结束符
    sentences = re.split(r'(?<=[。！？.!?])\s*', text)
    return [s.strip() for s in sentences if s.strip()]
