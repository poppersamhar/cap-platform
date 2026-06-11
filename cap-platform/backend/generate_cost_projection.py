"""生成 CAP 成本测算表（仅 Avatar 有缓存，含敏感性分析）

运行: python3 generate_cost_projection.py
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

C_DARK = "2D2A26"
C_WHITE = "FFFFFF"
C_GRAY = "8C857E"
C_NOTE = "6B655E"
C_RED = "C45C3E"
C_GREEN = "3A8A5E"
C_BEIGE = "E8E0D5"
C_HL = "FFF3E0"

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

PRICING = {"cache_read": 0.420, "standard_input": 2.100, "output": 8.400}
AVATAR_CACHE_RATIO = 0.40

# 基于 real_measurement_v2.py 实测数据（2026-06-11）
# 仅 Avatar 有缓存，其他 Agent 缓存命中率为 0%
# 分身工厂数据基于代码分析估算（100条记录→3个典型分身）
SCENARIOS = {
    "销售对练": {
        "desc": "Training mode，15轮对话 + 结束后1次督导评分 + 1次分析报告",
        "agents": [
            {"name": "avatar", "calls": 15, "prompt": 41540, "output": 1511, "cacheable": True},
            {"name": "supervisor", "calls": 1, "prompt": 2524, "output": 1850, "cacheable": False},
            {"name": "analyst", "calls": 1, "prompt": 878, "output": 1024, "cacheable": False},
        ],
        "monthly_calls": 500_000,
    },
    "用户调研": {
        "desc": "Research mode，20轮对话 + 1次洞察报告（无 Supervisor）",
        "agents": [
            {"name": "avatar", "calls": 20, "prompt": 45478, "output": 2249, "cacheable": True},
            {"name": "analyst", "calls": 1, "prompt": 1344, "output": 1024, "cacheable": False},
        ],
        "monthly_calls": 500_000,
    },
    "个体用户导入": {
        "desc": "Extractor Agent，单次调用提取结构化 Persona",
        "agents": [
            {"name": "extractor", "calls": 1, "prompt": 1369, "output": 538, "cacheable": False},
        ],
        "monthly_calls": 100_000,
    },
    "问卷调研": {
        "desc": "Survey Agent，10题开放式问卷 × 1个分身（实测 2026-06-11）",
        "agents": [
            {"name": "survey", "calls": 10, "prompt": 17435, "output": 3156, "cacheable": False, "note": "逐题调用，含完整画像+历史问答上下文（实测：杨波）"},
        ],
        "monthly_calls": 50_000,
    },
    "分身工厂": {
        "desc": "100条真实用户数据 → LLM特征抽取 → K-Means聚类 → LLM合成3个典型分身",
        "agents": [
            {"name": "factory_extractor", "calls": 100, "prompt": 220000, "output": 30000, "cacheable": False, "note": "逐条提取10维结构化特征（System≈800 + User≈1,400）×100"},
            {"name": "factory_synthesizer", "calls": 3, "prompt": 7800, "output": 9000, "cacheable": False, "note": "每簇合成1个典型分身（System≈900 + 统计摘要≈1,700）×3"},
        ],
        "monthly_calls": 100,  # 假设每月运行100次工厂
    },
}


def calc_cost(agents, pricing, avatar_cache_ratio=0.0):
    total_cost = 0
    for a in agents:
        prompt = a["prompt"]
        output = a["output"]
        cache_ratio = avatar_cache_ratio if a.get("cacheable") else 0.0
        cacheable = int(prompt * cache_ratio)
        noncache = prompt - cacheable
        cost = (
            cacheable / 1_000_000 * pricing["cache_read"]
            + noncache / 1_000_000 * pricing["standard_input"]
            + output / 1_000_000 * pricing["output"]
        )
        total_cost += cost
    return total_cost


def set_cell(ws, row, col, value, **kw):
    cell = ws.cell(row=row, column=col, value=value)
    cell.alignment = Alignment(
        horizontal=kw.get("align_h", "center"),
        vertical=kw.get("align_v", "center"),
        wrap_text=True,
    )
    cell.border = BORDER
    font_color = kw.get("font_color", C_DARK)
    if isinstance(font_color, tuple):
        font_color = font_color[0]
    cell.font = Font(
        name="微软雅黑",
        size=kw.get("font_size", 11),
        bold=kw.get("bold", False),
        color=font_color,
    )
    if "bg_color" in kw:
        cell.fill = PatternFill(start_color=kw["bg_color"], end_color=kw["bg_color"], fill_type="solid")
    if "num_fmt" in kw:
        cell.number_format = kw["num_fmt"]
    return cell


def build_sheet(ws):
    # ── 列宽 ──
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 16
    ws.column_dimensions["I"].width = 18

    row = 1

    # ═══ 大标题 ═══
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    set_cell(ws, row, 1, "CAP 平台 Token 成本测算表", font_size=18, bold=True, align_h="center", font_color=C_DARK)
    ws.row_dimensions[row].height = 35
    row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    info = "测算日期：2026-06-11    模型：MiniMax-M2.7    仅 Avatar 启用缓存（40%），其他 Agent 无缓存    定价来源：MiniMax 国内站"
    set_cell(ws, row, 1, info, font_size=10, align_h="center", font_color=C_GRAY)
    ws.row_dimensions[row].height = 22
    row += 2

    # ═══ 一、定价 ═══
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    set_cell(ws, row, 1, "一、MiniMax-M2.7 国内站定价", font_size=11, bold=True, align_h="left", font_color=C_DARK)
    row += 1

    for c, h in enumerate(["计费类型", "单价（¥/M tokens）", "说明"], 1):
        set_cell(ws, row, c, h, bg_color=C_DARK, font_color=C_WHITE, bold=True)
    ws.row_dimensions[row].height = 28
    row += 1

    pricing_rows = [
        ("Cache Read（缓存命中）", PRICING["cache_read"], "仅 Avatar 的 system prompt 前缀部分"),
        ("Standard Input（标准输入）", PRICING["standard_input"], "非缓存的 input tokens（所有 Agent）"),
        ("Output（输出）", PRICING["output"], "LLM 生成的 completion tokens（所有 Agent）"),
    ]
    for label, price, note in pricing_rows:
        is_total = label.startswith("Output")
        bg = C_BEIGE if is_total else None
        bold = is_total
        set_cell(ws, row, 1, label, bold=bold, bg_color=bg, align_h="left")
        set_cell(ws, row, 2, price, bold=bold, bg_color=bg, num_fmt='"¥"#,##0.000')
        set_cell(ws, row, 3, note, bold=bold, bg_color=bg, align_h="left")
        row += 1
    row += 1

    # ═══ 二、各场景明细 ═══
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    set_cell(ws, row, 1, "二、各场景 Token 消耗与成本明细（仅 Avatar 40% 缓存，其他 Agent 无缓存）", font_size=11, bold=True, align_h="left", font_color=C_DARK)
    row += 1

    for scene_name, scene_data in SCENARIOS.items():
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        set_cell(ws, row, 1, f"【{scene_name}】{scene_data['desc']}", font_size=11, bold=True, align_h="left", font_color="1A5276")
        ws.row_dimensions[row].height = 26
        row += 1

        agent_headers = ["Agent", "调用次数", "Input", "Cacheable", "Non-Cache", "Output", "单次成本（无缓存）", "单次成本（缓存后）", "月成本（缓存后）"]
        for c, h in enumerate(agent_headers, 1):
            set_cell(ws, row, c, h, bg_color=C_DARK, font_color=C_WHITE, bold=True)
        ws.row_dimensions[row].height = 26
        row += 1

        scene_cost_nc = 0
        scene_cost_ca = 0
        for a in scene_data["agents"]:
            prompt = a["prompt"]
            output = a["output"]
            cacheable = int(prompt * AVATAR_CACHE_RATIO) if a.get("cacheable") else 0
            noncache = prompt - cacheable

            # 无缓存单次成本
            cost_nc = prompt / 1_000_000 * PRICING["standard_input"] + output / 1_000_000 * PRICING["output"]
            # 缓存后单次成本
            cost_ca = cacheable / 1_000_000 * PRICING["cache_read"] + noncache / 1_000_000 * PRICING["standard_input"] + output / 1_000_000 * PRICING["output"]

            scene_cost_nc += cost_nc
            scene_cost_ca += cost_ca

            cache_note = f"{AVATAR_CACHE_RATIO*100:.0f}%" if a.get("cacheable") else "0%"
            set_cell(ws, row, 1, a["name"])
            set_cell(ws, row, 2, a["calls"])
            set_cell(ws, row, 3, prompt, num_fmt='#,##0')
            set_cell(ws, row, 4, cacheable, num_fmt='#,##0')
            set_cell(ws, row, 5, noncache, num_fmt='#,##0')
            set_cell(ws, row, 6, output, num_fmt='#,##0')
            set_cell(ws, row, 7, cost_nc, num_fmt='"¥"#,##0.0000', font_color=C_GRAY)
            set_cell(ws, row, 8, cost_ca, num_fmt='"¥"#,##0.0000', font_color=C_RED)
            set_cell(ws, row, 9, cost_ca * scene_data["monthly_calls"], num_fmt='"¥"#,##0.00', font_color=C_RED)
            row += 1

        # 场景小计
        set_cell(ws, row, 1, "场景小计", bold=True, bg_color=C_BEIGE, align_h="left")
        set_cell(ws, row, 2, "—", bold=True, bg_color=C_BEIGE)
        set_cell(ws, row, 3, "—", bold=True, bg_color=C_BEIGE)
        set_cell(ws, row, 4, "—", bold=True, bg_color=C_BEIGE)
        set_cell(ws, row, 5, "—", bold=True, bg_color=C_BEIGE)
        set_cell(ws, row, 6, "—", bold=True, bg_color=C_BEIGE)
        set_cell(ws, row, 7, scene_cost_nc, bold=True, bg_color=C_BEIGE, num_fmt='"¥"#,##0.0000', font_color=C_GRAY)
        set_cell(ws, row, 8, scene_cost_ca, bold=True, bg_color=C_BEIGE, num_fmt='"¥"#,##0.0000', font_color=C_RED)
        set_cell(ws, row, 9, scene_cost_ca * scene_data["monthly_calls"], bold=True, bg_color=C_BEIGE, num_fmt='"¥"#,##0.00', font_color=C_RED)
        row += 1

        # 说明
        yearly = scene_cost_ca * scene_data["monthly_calls"] * 12
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        note = f"  月调用量：{scene_data['monthly_calls']:,} 次  |  年成本（无缓存）：¥{scene_cost_nc * scene_data['monthly_calls'] * 12:,.2f}  |  年成本（缓存后）：¥{yearly:,.2f}"
        set_cell(ws, row, 1, note, font_size=9, align_h="left", font_color=C_GRAY, italic=True)
        row += 1
        row += 1

    # ═══ 三、总计 ═══
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    set_cell(ws, row, 1, "三、总计", font_size=11, bold=True, align_h="left", font_color=C_DARK)
    row += 1

    total_headers = ["场景", "月调用量（次）", "单次成本（无缓存）", "单次成本（缓存后）", "月成本（无缓存）", "月成本（缓存后）", "年成本（无缓存）", "年成本（缓存后）", "缓存节省"]
    for c, h in enumerate(total_headers, 1):
        set_cell(ws, row, c, h, bg_color=C_DARK, font_color=C_WHITE, bold=True)
    ws.row_dimensions[row].height = 28
    row += 1

    grand_nc = 0
    grand_ca = 0

    for scene_name in ["销售对练", "用户调研", "个体用户导入", "问卷调研", "分身工厂"]:
        data = SCENARIOS[scene_name]
        monthly_calls = data["monthly_calls"]
        cost_nc = calc_cost(data["agents"], PRICING, 0.0)
        cost_ca = calc_cost(data["agents"], PRICING, AVATAR_CACHE_RATIO)
        monthly_nc = cost_nc * monthly_calls
        monthly_ca = cost_ca * monthly_calls
        yearly_nc = monthly_nc * 12
        yearly_ca = monthly_ca * 12
        saving = yearly_nc - yearly_ca

        grand_nc += yearly_nc
        grand_ca += yearly_ca

        set_cell(ws, row, 1, scene_name, align_h="left")
        set_cell(ws, row, 2, monthly_calls, num_fmt='#,##0')
        set_cell(ws, row, 3, cost_nc, num_fmt='"¥"#,##0.0000', font_color=C_GRAY)
        set_cell(ws, row, 4, cost_ca, num_fmt='"¥"#,##0.0000', font_color=C_RED)
        set_cell(ws, row, 5, monthly_nc, num_fmt='"¥"#,##0.00', font_color=C_GRAY)
        set_cell(ws, row, 6, monthly_ca, num_fmt='"¥"#,##0.00', font_color=C_RED)
        set_cell(ws, row, 7, yearly_nc, num_fmt='"¥"#,##0.00', font_color=C_GRAY)
        set_cell(ws, row, 8, yearly_ca, num_fmt='"¥"#,##0.00', font_color=C_RED)
        set_cell(ws, row, 9, saving, num_fmt='"¥"#,##0.00', font_color=C_GREEN)
        row += 1

    total_saving = grand_nc - grand_ca
    set_cell(ws, row, 1, "合计", bold=True, bg_color=C_BEIGE, align_h="left")
    set_cell(ws, row, 2, "—", bold=True, bg_color=C_BEIGE)
    set_cell(ws, row, 3, "—", bold=True, bg_color=C_BEIGE)
    set_cell(ws, row, 4, "—", bold=True, bg_color=C_BEIGE)
    set_cell(ws, row, 5, "—", bold=True, bg_color=C_BEIGE)
    set_cell(ws, row, 6, "—", bold=True, bg_color=C_BEIGE)
    set_cell(ws, row, 7, grand_nc, bold=True, bg_color=C_BEIGE, num_fmt='"¥"#,##0.00', font_color=C_GRAY)
    set_cell(ws, row, 8, grand_ca, bold=True, bg_color=C_BEIGE, num_fmt='"¥"#,##0.00', font_color=C_RED)
    set_cell(ws, row, 9, total_saving, bold=True, bg_color=C_BEIGE, num_fmt='"¥"#,##0.00', font_color=C_GREEN)
    row += 2

    # ═══ 四、敏感性分析：Avatar 缓存命中率 ═══
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    set_cell(ws, row, 1, "四、敏感性分析：Avatar 缓存命中率对年成本的影响（仅 Avatar 有缓存，其他 Agent 无缓存）",
             font_size=11, bold=True, align_h="left", font_color=C_DARK)
    row += 1

    sens_headers = ["缓存命中率", "销售对练年成本", "用户调研年成本", "个体导入年成本", "问卷调研年成本", "分身工厂年成本", "合计年成本", "较无缓存节省", "节省比例"]
    for c, h in enumerate(sens_headers, 1):
        set_cell(ws, row, c, h, bg_color=C_DARK, font_color=C_WHITE, bold=True)
    ws.row_dimensions[row].height = 28
    row += 1

    # 计算无缓存基准
    base_costs = {}
    for scene_name in ["销售对练", "用户调研", "个体用户导入", "问卷调研", "分身工厂"]:
        data = SCENARIOS[scene_name]
        base_costs[scene_name] = calc_cost(data["agents"], PRICING, 0.0) * data["monthly_calls"] * 12

    base_total = sum(base_costs.values())

    for rate in [0.0, 0.30, 0.40, 0.50, 0.60, 0.70, 0.80, 0.90]:
        costs = {}
        for scene_name in ["销售对练", "用户调研", "个体用户导入"]:
            data = SCENARIOS[scene_name]
            costs[scene_name] = calc_cost(data["agents"], PRICING, rate) * data["monthly_calls"] * 12
        # 问卷调研 / 分身工厂不受 Avatar 缓存影响（其 Agent 非 Avatar）
        costs["问卷调研"] = base_costs["问卷调研"]
        costs["分身工厂"] = base_costs["分身工厂"]
        total = sum(costs.values())
        saving = base_total - total
        saving_pct = saving / base_total * 100 if base_total > 0 else 0

        is_highlight = abs(rate - AVATAR_CACHE_RATIO) < 0.01
        bg = C_HL if is_highlight else None
        bold = is_highlight

        set_cell(ws, row, 1, f"{rate*100:.0f}%", bold=bold, bg_color=bg)
        set_cell(ws, row, 2, costs["销售对练"], num_fmt='"¥"#,##0.00', bold=bold, bg_color=bg)
        set_cell(ws, row, 3, costs["用户调研"], num_fmt='"¥"#,##0.00', bold=bold, bg_color=bg)
        set_cell(ws, row, 4, costs["个体用户导入"], num_fmt='"¥"#,##0.00', bold=bold, bg_color=bg)
        set_cell(ws, row, 5, costs["问卷调研"], num_fmt='"¥"#,##0.00', bold=bold, bg_color=bg)
        set_cell(ws, row, 6, costs["分身工厂"], num_fmt='"¥"#,##0.00', bold=bold, bg_color=bg)
        set_cell(ws, row, 7, total, num_fmt='"¥"#,##0.00', bold=bold, bg_color=bg)
        set_cell(ws, row, 8, saving, num_fmt='"¥"#,##0.00', font_color=C_GREEN, bold=bold, bg_color=bg)
        set_cell(ws, row, 9, f"{saving_pct:.1f}%", font_color=C_GREEN, bold=bold, bg_color=bg)
        row += 1

    row += 1

    # ═══ 五、假设说明 ═══
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    set_cell(ws, row, 1, "五、测算假设与说明", font_size=11, bold=True, align_h="left", font_color=C_DARK)
    row += 1

    notes = [
        f"1. 定价来源：MiniMax 国内站（Cache Read ¥{PRICING['cache_read']}/M、Standard Input ¥{PRICING['standard_input']}/M、Output ¥{PRICING['output']}/M）",
        "2. 缓存策略：仅 Avatar Agent 的 system prompt 前缀部分可缓存（Persona JSON + 情绪指令 + 知识库上下文），缓存命中率按 Avatar 的 prompt 比例计算。",
        "3. Supervisor/Analyst/Extractor/Factory Agent 的 prompt 包含大量动态内容（完整对话历史、评估结果、单条记录等），每次调用差异大，基本不可缓存，统一按 0% 计算。",
        "4. 销售对练：15轮对话，avatar 41,540 prompt + 1,511 output；supervisor 结束后1次评估 2,524P + 1,850C；analyst 1次 878P + 1,024C。",
        "5. 用户调研：20轮对话，avatar 45,478 prompt + 2,249 output；analyst 1次 1,344P + 1,024C（research 模式无 supervisor）。",
        "6. 个体导入：单次 extractor 调用，1,369 prompt + 538 output。",
        "7. 问卷调研：Survey Agent 10次调用（每题1次），逐题含完整画像+历史问答，实测 17,435 prompt + 3,156 output（单分身，杨波）。",
        "8. 分身工厂：100条真实用户数据 → 特征抽取（100次，System≈800 + User≈1,400）×100 → K-Means聚类（本地计算，0 token）→ 合成典型分身（3次，System≈900 + 统计摘要≈1,700）×3。",
        "9. 注意：MiniMax Native API 未在 usage 中返回 cache_read_input_tokens 字段，实际缓存折扣是否生效需进一步确认。",
    ]
    for note in notes:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        set_cell(ws, row, 1, note, font_size=10, align_h="left", font_color=C_NOTE)
        ws.row_dimensions[row].height = 22
        row += 1


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "成本测算"
    build_sheet(ws)
    output_path = "/Users/samhar/Downloads/CAP_Cost_Projection.xlsx"
    wb.save(output_path)
    print(f"Excel 已生成: {output_path}")

    # 控制台概览
    print("\n" + "="*70)
    print("  CAP Token 成本测算（仅 Avatar 40% 缓存）")
    print("="*70)
    for name, data in SCENARIOS.items():
        cost_nc = calc_cost(data["agents"], PRICING, 0.0)
        cost_ca = calc_cost(data["agents"], PRICING, AVATAR_CACHE_RATIO)
        print(f"  {name:12s}: 单次(无缓存)=¥{cost_nc:.4f}  单次(缓存)=¥{cost_ca:.4f}  年(缓存)=¥{cost_ca*data['monthly_calls']*12:,.2f}")

    base_total = sum(calc_cost(d["agents"], PRICING, 0.0) * d["monthly_calls"] * 12 for d in SCENARIOS.values())
    ca_total = sum(calc_cost(d["agents"], PRICING, AVATAR_CACHE_RATIO) * d["monthly_calls"] * 12 for d in SCENARIOS.values())
    print(f"\n  {'合计':12s}: 年(无缓存)=¥{base_total:,.2f}  年(缓存)=¥{ca_total:,.2f}  节省=¥{base_total-ca_total:,.2f}（{(1-ca_total/base_total)*100:.1f}%）")

    print("\n  敏感性分析（年成本）：")
    for rate in [0.30, 0.40, 0.50, 0.60, 0.70, 0.80, 0.90]:
        total = sum(calc_cost(d["agents"], PRICING, rate) * d["monthly_calls"] * 12 for d in SCENARIOS.values())
        saving = base_total - total
        print(f"    Avatar 缓存 {rate*100:.0f}%: ¥{total:,.2f}  节省 ¥{saving:,.2f} ({saving/base_total*100:.1f}%)")


if __name__ == "__main__":
    main()
